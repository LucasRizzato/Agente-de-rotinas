"""
Planilha de controle (.xlsx) de notas fiscais — colaboradores PJ e
fornecedores compartilham a mesma planilha, diferenciados pela coluna
"Categoria".
"""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    "Categoria",
    "Mês",
    "Nome",
    "Número da Nota",
    "Mês de Emissão",
    "Data de Recebimento",
    "Arquivo NF",
    "Arquivo Boleto",
    "Status Pagamento",
    "Observações",
    "Assunto do E-mail",
    "Gmail Message ID",
]

CATEGORIA_PJ = "Colaborador PJ"
CATEGORIA_FORNECEDOR = "Fornecedor"


def _new_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Controle"
    ws.append(COLUMNS)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    for col_idx, _ in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
    ws.freeze_panes = "A2"
    for col_idx, header in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(header) + 4)
    return wb


def load_or_create(path: str) -> Workbook:
    sheet_path = Path(path)
    if sheet_path.exists():
        return load_workbook(sheet_path)
    sheet_path.parent.mkdir(parents=True, exist_ok=True)
    return _new_workbook()


def already_logged(wb: Workbook, gmail_message_id: str) -> bool:
    ws = wb["Controle"]
    id_col = COLUMNS.index("Gmail Message ID") + 1
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[id_col - 1].value == gmail_message_id:
            return True
    return False


def append_row(
    wb: Workbook,
    *,
    categoria: str,
    mes_pasta: str,
    nome: str,
    numero_nota: str,
    data_recebimento: datetime,
    mes_emissao: str = "",
    arquivo_nf: str = "",
    arquivo_boleto: str = "",
    observacoes: str = "",
    assunto: str = "",
    gmail_message_id: str = "",
):
    ws = wb["Controle"]
    ws.append(
        [
            categoria,
            mes_pasta,
            nome,
            numero_nota,
            mes_emissao,
            data_recebimento.strftime("%d/%m/%Y %H:%M"),
            arquivo_nf,
            arquivo_boleto,
            "Pendente",
            observacoes,
            assunto,
            gmail_message_id,
        ]
    )


def save(wb: Workbook, path: str):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
